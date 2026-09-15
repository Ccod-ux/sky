from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"D:\杨惠杰_Java实习简历_优化版.docx")
PHOTO = Path(r"C:\Users\ccod\Desktop\retouch_2024081218122959.jpg")

NAVY = "244A6B"
GRAY = "555555"
BLACK = "000000"
FONT = "Microsoft YaHei"


def set_run_font(run, size=9.2, bold=False, color=BLACK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    return run


def set_paragraph(paragraph, before=0, after=0, line=1.08, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_together = True
    if keep:
        fmt.keep_with_next = True
    return paragraph


def add_text(paragraph, text, size=9.2, bold=False, color=BLACK):
    return set_run_font(paragraph.add_run(text), size=size, bold=bold, color=color)


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn("w:" + margin))
        if element is None:
            element = OxmlElement("w:" + margin)
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_repeat_table_layout(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    remove_table_borders(table)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=0, start=0, bottom=0, end=0)


def add_section_title(doc, title):
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, before=4.8, after=1.8, line=1.0, keep=True)
    add_text(paragraph, title, size=11.8, bold=True, color=NAVY)
    return paragraph


def add_intent_item(cell, icon, text):
    paragraph = cell.paragraphs[0]
    set_paragraph(paragraph, after=0.8, line=1.0)
    add_text(paragraph, icon + "  ", size=9.3, bold=True, color=NAVY)
    add_text(paragraph, text, size=9.25, color=BLACK)


def add_project_heading(doc, date, title):
    table = doc.add_table(rows=1, cols=2)
    set_repeat_table_layout(table, [5.35, 12.75])
    left, right = table.rows[0].cells
    p = left.paragraphs[0]
    set_paragraph(p, after=0.5, line=1.0, keep=True)
    add_text(p, date, size=9.2, bold=True, color=BLACK)
    p = right.paragraphs[0]
    set_paragraph(p, after=0.5, line=1.0, keep=True)
    add_text(p, title, size=9.7, bold=True, color=BLACK)
    return table


def add_labeled_line(doc, label, text, size=8.75, after=0.7):
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, after=after, line=1.05)
    add_text(paragraph, label, size=size, bold=True, color=GRAY)
    add_text(paragraph, text, size=size, color=BLACK)
    return paragraph


def add_numbered_line(doc, number, text):
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, after=0.65, line=1.06)
    paragraph.paragraph_format.left_indent = Cm(0.52)
    paragraph.paragraph_format.first_line_indent = Cm(-0.52)
    add_text(paragraph, f"{number}. ", size=8.65, color=BLACK)
    add_text(paragraph, text, size=8.65, color=BLACK)
    return paragraph


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.05)
section.bottom_margin = Cm(1.05)
section.left_margin = Cm(1.18)
section.right_margin = Cm(1.18)
section.header_distance = Cm(0.5)
section.footer_distance = Cm(0.5)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(9.2)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.08
normal_rpr = normal._element.get_or_add_rPr()
normal_rfonts = normal_rpr.rFonts
if normal_rfonts is None:
    normal_rfonts = OxmlElement("w:rFonts")
    normal_rpr.insert(0, normal_rfonts)
normal_rfonts.set(qn("w:ascii"), FONT)
normal_rfonts.set(qn("w:hAnsi"), FONT)
normal_rfonts.set(qn("w:eastAsia"), FONT)

title_style = doc.styles["Title"]
title_style.font.name = FONT
title_style.font.size = Pt(19.5)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
title_style.paragraph_format.space_before = Pt(0)
title_style.paragraph_format.space_after = Pt(4)
title_style.paragraph_format.line_spacing = 1.0
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
title_rpr = title_style._element.get_or_add_rPr()
title_rfonts = title_rpr.rFonts
if title_rfonts is None:
    title_rfonts = OxmlElement("w:rFonts")
    title_rpr.insert(0, title_rfonts)
title_rfonts.set(qn("w:ascii"), FONT)
title_rfonts.set(qn("w:hAnsi"), FONT)
title_rfonts.set(qn("w:eastAsia"), FONT)

# Top block follows the supplied PDF: title and personal details on the left,
# portrait aligned at the upper-right edge.
header = doc.add_table(rows=1, cols=2)
set_repeat_table_layout(header, [15.85, 2.75])
left, right = header.rows[0].cells
left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

p = left.paragraphs[0]
p.style = doc.styles["Title"]
set_paragraph(p, after=4, line=1.0, keep=True)
add_text(p, "杨惠杰的简历", size=19.5, bold=True, color=BLACK)

p = left.add_paragraph()
set_paragraph(p, after=3.5, line=1.0, keep=True)
add_text(p, "Java 后端开发方向，具备 Spring Boot、MySQL、Redis 项目实践。", size=9.1, color=BLACK)

p = left.add_paragraph()
set_paragraph(p, after=2.1, line=1.0)
p.paragraph_format.tab_stops.add_tab_stop(Cm(6.15), WD_TAB_ALIGNMENT.LEFT)
p.paragraph_format.tab_stops.add_tab_stop(Cm(11.55), WD_TAB_ALIGNMENT.LEFT)
add_text(p, "男", size=9.2)
add_text(p, "\t本科", size=9.2)
add_text(p, "\t20 岁", size=9.2)

p = left.add_paragraph()
set_paragraph(p, after=2.1, line=1.0)
p.paragraph_format.tab_stops.add_tab_stop(Cm(6.15), WD_TAB_ALIGNMENT.LEFT)
p.paragraph_format.tab_stops.add_tab_stop(Cm(11.55), WD_TAB_ALIGNMENT.LEFT)
add_text(p, "杭州", size=9.2)
add_text(p, "\t13362055279", size=9.2)

p = left.add_paragraph()
set_paragraph(p, after=0, line=1.0)
add_text(p, "2871197445@qq.com", size=9.2)

p = right.paragraphs[0]
set_paragraph(p, after=0, line=1.0)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
picture = p.add_run().add_picture(str(PHOTO), width=Cm(2.3))
picture._inline.docPr.set("descr", "杨惠杰证件照")

add_section_title(doc, "求职意向")
intent = doc.add_table(rows=2, cols=2)
set_repeat_table_layout(intent, [9.3, 9.3])
add_intent_item(intent.cell(0, 0), "★", "Java 后端开发实习生")
add_intent_item(intent.cell(0, 1), "★", "杭州")
add_intent_item(intent.cell(1, 0), "●", "薪资面议")
add_intent_item(intent.cell(1, 1), "★", "一个月内可到岗")

add_section_title(doc, "教育背景")
education = doc.add_table(rows=1, cols=2)
set_repeat_table_layout(education, [5.35, 12.75])
p = education.cell(0, 0).paragraphs[0]
set_paragraph(p, after=0.8, line=1.0, keep=True)
add_text(p, "2024.09 - 2028.06", size=9.25, bold=True)
p = education.cell(0, 1).paragraphs[0]
set_paragraph(p, after=0.8, line=1.0, keep=True)
add_text(p, "浙江农林大学", size=9.7, bold=True)
p = doc.add_paragraph()
set_paragraph(p, after=0.8, line=1.0, keep=True)
add_text(p, "数据科学与大数据技术专业 | 本科", size=9.25, bold=True)
add_labeled_line(
    doc,
    "主修课程：",
    "数据结构与算法、数据库原理、Java 程序设计、计算机网络、操作系统",
    size=8.7,
    after=0,
)

add_section_title(doc, "项目经历")
add_project_heading(doc, "2026.03 - 2026.06", "校区订餐系统 | Java 后端开发")
add_labeled_line(
    doc,
    "项目简介：",
    "面向高校场景的在线点餐系统，支持菜品管理、套餐与优惠券秒杀、用户下单、支付及催单。",
    size=8.65,
    after=0.45,
)
add_labeled_line(
    doc,
    "技术栈：",
    "Spring Boot、MySQL、MyBatis、Redis、RabbitMQ、阿里云 OSS、JWT",
    size=8.55,
    after=0.45,
)
add_numbered_line(doc, 1, "基于 JWT 与拦截器实现无状态认证和统一身份校验，并通过 ThreadLocal 传递当前用户上下文。")
add_numbered_line(doc, 2, "使用 Redis + Lua 完成库存与一人一单校验，防止超卖；通过防重 Token + Lua 保证下单幂等。")
add_numbered_line(doc, 3, "将小程序高频查询数据缓存至 Redis，并结合布隆过滤器降低缓存穿透风险。")
add_numbered_line(doc, 4, "通过 RabbitMQ 异步创建秒杀订单，并发送延迟消息检查订单支付状态。")

add_project_heading(doc, "2025.11 - 2026.02", "广大农产购 | Java 后端开发")
add_labeled_line(
    doc,
    "项目简介：",
    "面向学校食堂食材采购的购物平台，包含运营管理后台与移动端，支持农产品浏览、购物车及下单。",
    size=8.65,
    after=0.45,
)
add_labeled_line(
    doc,
    "技术栈：",
    "Spring Boot、MyBatis-Plus、MySQL、Redis、Spring Task、RabbitMQ、WebSocket、JWT、OSS",
    size=8.55,
    after=0.45,
)
add_numbered_line(doc, 1, "使用 Redis 缓存商家营业状态、农产品及购物车数据；通过 JWT、拦截器与 ThreadLocal 处理用户认证。")
add_numbered_line(doc, 2, "使用 Spring Task 自动关闭超过 15 分钟未支付的订单，并将异常派送订单发送至 MQ 交由客服跟进。")
add_numbered_line(doc, 3, "通过 WebSocket 推送订单消息，使用 AOP 自动填充公共字段，并将图片、音视频资源存储至阿里云 OSS。")

add_section_title(doc, "专业技能")
add_labeled_line(
    doc,
    "Java 与框架：",
    "熟悉面向对象、集合、多线程和 IO；使用 Spring Boot、SSM、MyBatis / MyBatis-Plus 开发 Web 项目。",
    size=8.65,
    after=0.7,
)
add_labeled_line(
    doc,
    "数据库与中间件：",
    "熟悉 MySQL、SQL、索引及事务；掌握 Redis 缓存与 Lua 原子操作，了解 RabbitMQ 异步及延迟消息。",
    size=8.65,
    after=0.7,
)
add_labeled_line(
    doc,
    "开发工具：",
    "熟悉 Maven、Git，了解 Linux 常用操作；有 JWT、Spring Task、WebSocket 与 OSS 项目实践。",
    size=8.65,
    after=0,
)

doc.core_properties.title = "杨惠杰 Java 后端开发实习简历"
doc.core_properties.subject = "Java 后端开发实习求职简历"
doc.core_properties.author = "杨惠杰"
doc.save(OUTPUT)
print(OUTPUT)
