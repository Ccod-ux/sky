import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


INPUT = Path(r"C:\Users\ccod\Desktop\杨惠杰_Java实习简历_优化版.docx")
OUTPUT = Path(r"C:\Users\ccod\Desktop\杨惠杰_Java实习简历_大字版.docx")

SECTION_HEADINGS = {"求职意向", "教育背景", "项目经历", "专业技能"}
LABELED_PREFIXES = (
    "主修课程：",
    "项目简介：",
    "技术栈：",
    "Java 与框架：",
    "数据库与中间件：",
    "开发工具：",
)


def iter_paragraphs(container):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def remove_paragraph_border(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def remove_style_border(style):
    if style.type != WD_STYLE_TYPE.PARAGRAPH:
        return
    ppr = style._element.find(qn("w:pPr"))
    if ppr is None:
        return
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def force_nil_border(parent_properties, border_tag, edges):
    borders = parent_properties.find(qn("w:" + border_tag))
    if borders is None:
        borders = OxmlElement("w:" + border_tag)
        parent_properties.append(borders)
    for edge in edges:
        tag = qn("w:" + edge)
        element = borders.find(tag)
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "nil")
        element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "FFFFFF")


def remove_all_table_lines(table):
    force_nil_border(
        table._tbl.tblPr,
        "tblBorders",
        ("top", "left", "bottom", "right", "insideH", "insideV"),
    )
    for row in table.rows:
        for cell in row.cells:
            force_nil_border(
                cell._tc.get_or_add_tcPr(),
                "tcBorders",
                ("top", "left", "bottom", "right", "insideH", "insideV"),
            )
            for nested in cell.tables:
                remove_all_table_lines(nested)


def enlarged_size(old_size):
    if old_size is None:
        return None
    points = old_size.pt
    if points <= 8.6:
        return Pt(9.2)
    if points <= 9.05:
        return Pt(9.6)
    if points <= 9.35:
        return Pt(9.8)
    if points <= 9.7:
        return Pt(10.15)
    if points <= 11.7:
        return Pt(12.3)
    if points <= 20:
        return Pt(20.7)
    return Pt(points + 0.4)


def tune_paragraph(paragraph):
    text = paragraph.text.strip()
    fmt = paragraph.paragraph_format
    remove_paragraph_border(paragraph)

    for run in paragraph.runs:
        if run.text:
            new_size = enlarged_size(run.font.size)
            if new_size is not None:
                run.font.size = new_size

    if not text:
        return

    fmt.keep_together = True

    if text == "杨惠杰的简历":
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(5.5)
        fmt.line_spacing = 1.0
    elif text.startswith("Java 后端开发方向"):
        fmt.space_after = Pt(6.5)
        fmt.line_spacing = 1.05
    elif text in SECTION_HEADINGS:
        fmt.space_before = Pt(9.0)
        fmt.space_after = Pt(4.0)
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
    elif text.startswith(("男", "杭州", "2871197445")):
        fmt.space_after = Pt(3.0 if not text.startswith("2871197445") else 0)
        fmt.line_spacing = 1.0
    elif text.startswith(("★", "●")):
        fmt.space_after = Pt(2.2)
        fmt.line_spacing = 1.0
    elif text.startswith("数据科学与大数据技术专业"):
        fmt.space_after = Pt(1.3)
        fmt.line_spacing = 1.04
    elif text.startswith(LABELED_PREFIXES):
        fmt.space_after = Pt(1.8)
        fmt.line_spacing = 1.11
    elif re.match(r"^[1-4]\.\s", text):
        fmt.space_after = Pt(2.2)
        fmt.line_spacing = 1.12
    elif re.match(r"^20\d{2}\.\d{2}\s*-\s*20\d{2}\.\d{2}$", text):
        if text.startswith(("2025", "2026")):
            fmt.space_before = Pt(4.0)
        fmt.space_after = Pt(1.8)
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
    elif "| Java 后端开发" in text or text == "浙江农林大学":
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(1.2)
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
    else:
        fmt.space_after = Pt(1.0)
        fmt.line_spacing = 1.08


document = Document(INPUT)

document.styles["Normal"].font.size = Pt(9.2)
document.styles["Title"].font.size = Pt(20.7)
for style in document.styles:
    remove_style_border(style)

for paragraph in iter_paragraphs(document):
    tune_paragraph(paragraph)

for table in document.tables:
    remove_all_table_lines(table)

document.core_properties.title = "杨惠杰 Java 后端开发实习简历"
document.core_properties.subject = "Java 后端开发实习求职简历"
document.core_properties.author = "杨惠杰"
document.save(OUTPUT)
print(OUTPUT)
